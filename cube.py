import os
import time
from datetime import datetime
import openpyxl
import win32com.client as win32

# Define paths (use absolute raw paths for mapped drive locations)
base_dir = os.path.abspath(".")
input_param_file = r"Z:\CREDITORS\1. SI Recons\RPA - AP Payment Posted To Bank BOT\InPut_Parameter_File.xlsx"
template_file = os.path.join(base_dir, "Supplier Payments & Remittance download.xlsx")
base_output_dir = r"Z:\CREDITORS\1. SI Recons\RPA - AP Payment Posted To Bank BOT\Processed_Payments"

# Create a dated subfolder with '-cube' suffix (e.g., '2026-10-02-cube')
current_date_str = datetime.now().strftime("%Y-%m-%d")
output_dir = os.path.join(base_output_dir, f"{current_date_str}-cube")

os.makedirs(output_dir, exist_ok=True)
print(f"Output files will be saved to: {output_dir}")

# 1. Read Vendor Names (Col B) and PayNumbers (Col C) from InPut_Parameter_File.xlsx
print("Reading Vendor Names and PayNumbers from input file...")
wb_param = openpyxl.load_workbook(input_param_file, data_only=True)
ws_account = wb_param["Account_Details"]

records = []
for row in range(2, ws_account.max_row + 1):
    vendor_val = ws_account.cell(row=row, column=2).value
    pay_val = ws_account.cell(row=row, column=3).value

    if pay_val is not None and str(pay_val).strip():
        vendor_name = str(vendor_val).strip() if vendor_val is not None else "Unknown_Vendor"
        pay_number = str(pay_val).strip()
        records.append({"vendor": vendor_name, "pay_num": pay_number})

print(f"Found {len(records)} records to process.")

# 2. Launch a fresh isolated Excel instance via DispatchEx
excel = win32.DispatchEx("Excel.Application")

# Allow COM bridge to stabilize and set properties safely
for _ in range(5):
    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        break
    except AttributeError:
        time.sleep(1)

try:
    for idx, item in enumerate(records, start=1):
        vendor_name = item["vendor"]
        pay_num = item["pay_num"]
        print(f"\n[{idx}/{len(records)}] Processing Vendor: '{vendor_name}' | Payment Number: {pay_num}")

        # Open the template workbook
        wb = excel.Workbooks.Open(template_file)
        ws = wb.Sheets(1)
        pt = ws.PivotTables(1)

        # Locate the OLAP Company and Payment Number PivotFields
        payment_field = None
        company_field = None
        
        for pf in pt.PageFields:
            if "payment number" in pf.Name.lower():
                payment_field = pf
            elif "company" in pf.Name.lower():
                company_field = pf

        if payment_field is None:
            raise ValueError("Could not locate 'Payment Number' PageField in PivotTable.")

        # --- Set Company Filter based on Payment Number Prefix ---
        target_company = None
        if pay_num.startswith("TM"):
            target_company = "T&M Company owned stores"
        elif pay_num.startswith("TW"):
            target_company = "TWT Company owned stores"
            
        if target_company and company_field is not None:
            olap_company_name = f"[Supplier Payments].[Company].&[{target_company}]"
            company_filter_set = False
            try:
                company_field.CurrentPageName = olap_company_name
                company_filter_set = True
            except Exception:
                try:
                    company_field.CurrentPage = olap_company_name
                    company_filter_set = True
                except Exception:
                    try:
                        # Fallback to plain caption
                        company_field.CurrentPageName = target_company
                        company_filter_set = True
                    except Exception as err:
                        print(f"  Error setting company filter for {target_company}: {err}")
            
            if company_filter_set:
                print(f"  Applied Company filter: {target_company}")

        # --- Set Payment Number Filter ---
        olap_unique_name = f"[Supplier Payments].[Payment Number].&[{pay_num}]"
        filter_set = False
        try:
            payment_field.CurrentPageName = olap_unique_name
            filter_set = True
        except Exception:
            try:
                payment_field.CurrentPage = olap_unique_name
                filter_set = True
            except Exception:
                try:
                    # Fallback to plain caption
                    payment_field.CurrentPageName = pay_num
                    filter_set = True
                except Exception as final_err:
                    print(f"  Error setting payment filter for {pay_num}: {final_err}")

        if filter_set:
            print(f"  Applied Payment filter: {olap_unique_name}")

        # Wait for OLAP background query to finish retrieving data
        print("  Waiting for OLAP background query to complete...")
        wb.RefreshAll()
        excel.CalculateUntilAsyncQueriesDone()
        time.sleep(2)

        # Sanitize vendor and payment number for Windows filenames
        clean_vendor = "".join(c for c in vendor_name if c.isalnum() or c in ("-", "_", " ")).strip()
        clean_pay = "".join(c for c in pay_num if c.isalnum() or c in ("-", "_")).strip()

        # Generate filename incorporating Vendor Name and Payment Number
        file_name = f"{clean_vendor}_{clean_pay}.xlsx"
        save_path = os.path.join(output_dir, file_name)

        wb.SaveAs(save_path)
        wb.Close(SaveChanges=False)
        print(f"  Saved successfully: {save_path}")

    print("\nAll records processed successfully!")

finally:
    if excel is not None:
        try:
            excel.Quit()
        except Exception:
            pass